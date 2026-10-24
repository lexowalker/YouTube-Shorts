"""Keep docs/release-schedule.xlsx current and check whether the Shorts are on track for their release dates.

  python3 -I pipeline/schedule.py mark <chapter> <type> <status> [--title "…"] [--note "…"] [--created YYYY-MM-DD]
      Record one Short. type: reading | d1 | d2 (Discussion 1 / Discussion 2). status: Planned | "In progress" |
      Ready | Scheduled | Released. Moving to Ready or later fills "Created on" with today unless it's set already.
  python3 -I pipeline/schedule.py status [--today YYYY-MM-DD]
      On-track report: overdue Shorts, what's due in the next two weeks and whether it's made, and the buffer
      (how many upcoming Shorts in a row are already made).

The workbook keeps its formulas: release dates are computed in it from the Settings tab, and this script works out the
same dates from those inputs, so it never needs a recalculation to read them. Only Status, Short title, Notes and
Created on are written.
"""
import argparse, datetime, os, sys
from openpyxl import load_workbook

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BOOK = os.path.join(REPO, "docs", "release-schedule.xlsx")
TYPES = {"reading": "Reading", "d1": "Discussion 1", "d2": "Discussion 2"}
STATUSES = ["Planned", "In progress", "Ready", "Scheduled", "Released"]
MADE = {"Ready", "Scheduled", "Released"}
COL = {"chapter": 3, "type": 5, "date": 7, "status": 10, "title": 11, "notes": 12, "created": 13}


def as_date(v):
    return v.date() if isinstance(v, datetime.datetime) else v


def rows(wb):
    """Every Short with its release date, worked out from the Settings inputs exactly as the sheet's formula does."""
    st, sh = wb["Settings"], wb["Schedule"]
    start, step = as_date(st["B4"].value), int(st["B11"].value)
    offset = {st.cell(r, 1).value: int(st.cell(r, 2).value) for r in range(7, 10)}
    out = []
    for r in range(2, sh.max_row + 1):
        n = sh.cell(r, 1).value
        if n is None:
            continue
        typ = sh.cell(r, COL["type"]).value
        week = (n - 1) // 3 + 1
        out.append({"row": r, "n": n, "chapter": sh.cell(r, COL["chapter"]).value, "type": typ,
                    "date": start + datetime.timedelta(days=(week - 1) * step + offset[typ]),
                    "status": sh.cell(r, COL["status"]).value or "Planned", "title": sh.cell(r, COL["title"]).value})
    return out


def save(wb):
    wb.calculation.fullCalcOnLoad = True  # Excel recomputes every formula when the file is opened
    wb.save(BOOK)


def cmd_mark(a):
    wb = load_workbook(BOOK)
    sh = wb["Schedule"]
    typ = TYPES[a.type.lower()]
    hit = [r for r in rows(wb) if r["chapter"] == a.chapter and r["type"] == typ]
    if not hit:
        sys.exit(f"no row for chapter {a.chapter} {typ}")
    r = hit[0]["row"]
    status = next((s for s in STATUSES if s.lower() == a.status.lower()), None)
    if not status:
        sys.exit(f"status must be one of {', '.join(STATUSES)}")
    sh.cell(r, COL["status"]).value = status
    if a.title:
        sh.cell(r, COL["title"]).value = a.title
    if a.note:
        old = sh.cell(r, COL["notes"]).value
        sh.cell(r, COL["notes"]).value = f"{old}; {a.note}" if old else a.note
    created = sh.cell(r, COL["created"])
    if a.created:
        created.value = datetime.date.fromisoformat(a.created)
    elif status in MADE and not created.value:
        created.value = datetime.date.today()
    created.number_format = "d mmm yyyy"
    save(wb)
    print(f"#{hit[0]['n']} chapter {a.chapter} {typ} (release {hit[0]['date']:%a %d %b %Y}): {status}")


def cmd_status(a):
    today = datetime.date.fromisoformat(a.today) if a.today else datetime.date.today()
    rs = rows(load_workbook(BOOK))
    released = [r for r in rs if r["status"] == "Released"]
    overdue = [r for r in rs if r["date"] < today and r["status"] != "Released"]
    upcoming = [r for r in rs if r["date"] >= today and r["status"] != "Released"]
    fmt = lambda r: f"#{r['n']:<3} {r['date']:%a %d %b %Y}  Ch {r['chapter']:<2} {r['type']:<13} {r['status']}"
    print(f"today {today:%a %d %b %Y}: {len(released)} of {len(rs)} released")
    if overdue:
        print(f"\nBEHIND: {len(overdue)} past their release date and not released")
        for r in overdue:
            print("  " + fmt(r))
    soon = [r for r in upcoming if (r["date"] - today).days <= 14]
    if soon:
        print("\nnext two weeks:")
        for r in soon:
            print("  " + fmt(r) + ("" if r["status"] in MADE else "   <- not made yet"))
    # buffer: upcoming Shorts in release order that are already made, without a gap
    buf = 0
    for r in upcoming:
        if r["status"] not in MADE:
            break
        buf += 1
    nxt = upcoming[buf] if buf < len(upcoming) else None
    print(f"\nbuffer: the next {buf} Short(s) are made" +
          (f"; first not made is #{nxt['n']} (Ch {nxt['chapter']} {nxt['type']}), due {nxt['date']:%a %d %b}, "
           f"{(nxt['date'] - today).days} day(s) from today" if nxt else "; everything is made"))
    late = [r for r in soon if r["status"] not in MADE and (r["date"] - today).days <= 2]
    verdict = "BEHIND" if overdue else "AT RISK" if late else "ON TRACK"
    print(f"\n{verdict}" + (": a Short due within 2 days isn't made yet" if verdict == "AT RISK" else ""))


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    m = sub.add_parser("mark")
    m.add_argument("chapter", type=int)
    m.add_argument("type", choices=sorted(TYPES))
    m.add_argument("status")
    m.add_argument("--title")
    m.add_argument("--note")
    m.add_argument("--created")
    s = sub.add_parser("status")
    s.add_argument("--today")
    a = ap.parse_args()
    {"mark": cmd_mark, "status": cmd_status}[a.cmd](a)


if __name__ == "__main__":
    main()
